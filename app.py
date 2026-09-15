import streamlit as st
import pandas as pd
import time
import os
import re
import json
import hashlib
import unicodedata
from io import BytesIO
from pydantic import BaseModel, Field
from google import genai
from google.genai import types
from pypdf import PdfReader
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# 1. Estrutura de Extração
class ItemLicitacao(BaseModel):
    numero_item: str = Field(description="Número ou identificador do item/lote no edital (ex: '01', 'Lote 1 - Item 2')")
    descricao: str = Field(description="Descrição completa do medicamento no edital (princípio ativo, dosagem, forma)")
    principio_ativo_identificado: str = Field(description="Substância / Denominação genérica do princípio ativo identificado")
    unidade: str = Field(description="Unidade de medida/fornecimento (ex: AMP, FA, COMP, FR, BISNAGA)")
    quantidade: float = Field(description="Quantidade demandada expressa em número float")
    valor_referencia_unitario: float = Field(description="Valor unitário máximo ou de referência do edital")

class ListaItens(BaseModel):
    itens: list[ItemLicitacao]

# 2. Configurações da Página
st.set_page_config(page_title="Radar Farma - Licitações", layout="wide", page_icon="💊")
st.title("🎯 Analisador de Editais & Mapa de Preços")
st.caption("🚀 Pipeline Ativo: v3.3 - Cache em Dupla Camada (Instantâneo) & Barra de Progresso")

# Chave de API higienizada
api_key = ""
if "GEMINI_API_KEY" in st.secrets and st.secrets["GEMINI_API_KEY"]:
    api_key = str(st.secrets["GEMINI_API_KEY"]).strip().strip("'").strip('"')
else:
    api_key_input = st.sidebar.text_input("Chave Gemini API", type="password")
    if api_key_input:
        api_key = api_key_input.strip().strip("'").strip('"')

# 3. Base de Dados, Memória e Normalização
ARQUIVO_PORTFOLIO = "portfolio_laboratorios.xlsx"
ARQUIVO_MEMORIA = "memoria_medicamentos.json"

def normalizar_texto(texto):
    if not texto or pd.isna(texto):
        return ""
    nfkd = unicodedata.normalize('NFKD', str(texto))
    texto_sem_acento = "".join([c for c in nfkd if not unicodedata.combining(c)])
    texto_limpo = re.sub(r'[^A-Z0-9\s]', ' ', texto_sem_acento.upper())
    return " ".join(texto_limpo.split())

def carregar_memoria():
    if os.path.exists(ARQUIVO_MEMORIA):
        try:
            with open(ARQUIVO_MEMORIA, "r", encoding="utf-8") as f:
                return json.load(f)
        except Exception:
            return {}
    return {}

def salvar_aprendizado(novos_itens):
    memoria = carregar_memoria()
    atualizado = False
    for item in novos_itens:
        desc_norm = normalizar_texto(item.get("descricao", ""))
        princ_ativo = normalizar_texto(item.get("principio_ativo_identificado", ""))
        unidade = item.get("unidade", "")
        if desc_norm and princ_ativo and desc_norm not in memoria:
            memoria[desc_norm] = {
                "principio_ativo": princ_ativo,
                "unidade": unidade,
                "vezes_visto": 1,
                "ultima_atualizacao": time.strftime("%Y-%m-%d %H:%M:%S")
            }
            atualizado = True
    if atualizado:
        try:
            with open(ARQUIVO_MEMORIA, "w", encoding="utf-8") as f:
                json.dump(memoria, f, ensure_ascii=False, indent=2)
        except Exception:
            pass

MAPEAMENTO_LABS = {
    "Blau": ["BLAU"],
    "Eurofarma": ["EUROFARMA"],
    "Baxter": ["BAXTER"],
    "Biocon": ["BIOCON"],
    "Accord": ["ACCORD"],
    "Halex Istar": ["HALEX", "ISTAR"],
    "United Medical": ["UNITED MEDICAL", "UNITED"],
    "GSK": ["GSK", "GLAXO", "GLAXOSMITHKLINE"],
    "Aspen": ["ASPEN"],
    "Sanofi": ["SANOFI"],
    "Pint Pharma": ["PINT", "PINT PHARMA"]
}

@st.cache_data
def carregar_base_portfolio(caminho):
    if not os.path.exists(caminho):
        return None
    try:
        df = pd.read_excel(caminho)
        df.columns = [str(c).strip().upper() for c in df.columns]
        for col in ["SUBSTÂNCIA", "LABORATÓRIO", "PRODUTO", "APRESENTAÇÃO"]:
            if col in df.columns:
                df[col] = df[col].astype(str).str.strip()
        return df
    except Exception:
        return None

df_portfolio = carregar_base_portfolio(ARQUIVO_PORTFOLIO)
base_memoria = carregar_memoria()

# 4. Barra Lateral
st.sidebar.header("⚙️ Configurações de Busca")

segmento = st.sidebar.selectbox(
    "Segmento do Edital:",
    options=["Medicamentos", "Material Elétrico / Engenharia", "Personalizado / Geral"]
)

if segmento == "Medicamentos":
    st.sidebar.subheader("💊 Portfólio de Laboratórios")
    labs_base = [
        "Blau", "Eurofarma", "Baxter", "Biocon", "Accord",
        "Halex Istar", "United Medical", "GSK", "Aspen", "Sanofi", "Pint Pharma"
    ]
    
    labs_selecionados = st.sidebar.multiselect(
        "Laboratórios ativos para cotação:",
        options=labs_base,
        default=labs_base
    )
    
    if df_portfolio is not None:
        st.sidebar.success(f"Base Portfólio: {len(df_portfolio)} produtos cadastrados.")
    else:
        st.sidebar.warning("Arquivo 'portfolio_laboratorios.xlsx' não encontrado na raiz.")

    st.sidebar.info(f"🧠 **Memória de Aprendizado:** {len(base_memoria)} termos memorizados.")

    filtro_exibicao = st.sidebar.radio(
        "Visualização dos Resultados:",
        ["Todos os Medicamentos do Edital", "Apenas Itens com Match nos Laboratórios"],
        index=0
    )

    st.sidebar.markdown(
        "**Hierarquia Comercial:**\n"
        "- 🥇 **Sanofi** prioridade máxima.\n"
        "- 🥈 **Blau** sobre **Eurofarma**."
    )

elif segmento == "Material Elétrico / Engenharia":
    portfolio_texto = st.sidebar.text_area(
        "Itens de interesse:",
        value="Material elétrico, cabeamento estruturado, iluminação LED, quadros de distribuição"
    )
else:
    portfolio_texto = st.sidebar.text_area("Itens de interesse:", value="")

arquivo_pdf = st.file_uploader("Arraste o PDF do Edital ou Termo de Referência aqui", type=["pdf"])

# ========================================================
# PRÉ-PROCESSADOR CACHEADO (EXECUTA 1 VEZ POR PDF)
# ========================================================
@st.cache_data(show_spinner=False)
def pre_processar_pdf_inteligente(bytes_arquivo):
    leitor = PdfReader(BytesIO(bytes_arquivo))
    total_paginas = len(leitor.pages)
    
    termos_itens = [
        "ITEM", "QUANTIDADE", "VALOR ESTIMADO", "VALOR DE REFERENCIA", 
        "PRINCIPIO ATIVO", "FORMA FARMACEUTICA", "CONCENTRACAO", "LOTE",
        "TERMO DE REFERENCIA", "ANEXO", "ESPECIFICACAO"
    ]
    
    paginas_selecionadas = []
    paginas_indices = []

    for i, pagina in enumerate(leitor.pages):
        raw_text = pagina.extract_text() or ""
        norm_text = normalizar_texto(raw_text)
        score = sum(1 for t in termos_itens if t in norm_text)
        
        if score >= 2 or re.search(r'\b(ITEM\s+\d+|LOTE\s+\d+)\b', norm_text):
            paginas_selecionadas.append(f"--- PÁGINA {i+1} ---\n{raw_text}")
            paginas_indices.append(i + 1)

    if len(paginas_selecionadas) < 4:
        paginas_selecionadas = []
        paginas_indices = list(range(1, total_paginas + 1))
        for i, pagina in enumerate(leitor.pages):
            txt = pagina.extract_text() or ""
            paginas_selecionadas.append(f"--- PÁGINA {i+1} ---\n{txt}")

    texto_filtrado = "\n\n".join(paginas_selecionadas)
    return texto_filtrado, paginas_indices, total_paginas

# ========================================================
# FUNÇÃO DE EXTRAÇÃO COM CACHE (RETORNO IMEDIATO)
# ========================================================
@st.cache_data(show_spinner=False)
def extrair_itens_com_gemini_cached(pdf_bytes_hash, texto_edital, prompt_instrucao, key_api):
    client = genai.Client(api_key=key_api)
    modelo_eleito = "gemini-3.6-flash"

    try:
        modelos_ativos = [m.name.replace("models/", "") for m in client.models.list()]
        flashes_validos = [m for m in modelos_ativos if "flash" in m.lower() and "embed" not in m.lower()]
        if "gemini-3.6-flash" in flashes_validos:
            modelo_eleito = "gemini-3.6-flash"
        elif flashes_validos:
            modelo_eleito = flashes_validos[0]
    except Exception:
        modelo_eleito = "gemini-3.6-flash"

    resposta = None
    ultimo_erro = None
    max_tentativas = 4

    for tentativa in range(1, max_tentativas + 1):
        try:
            resposta = client.models.generate_content(
                model=modelo_eleito,
                contents=prompt_instrucao,
                config=types.GenerateContentConfig(
                    response_mime_type="application/json",
                    response_schema=ListaItens,
                )
            )
            if resposta and resposta.text:
                break
        except Exception as err:
            ultimo_erro = err
            msg_erro = str(err)
            if any(c in msg_erro for c in ["429", "RESOURCE_EXHAUSTED", "503", "UNAVAILABLE"]):
                match_tempo = re.search(r'retry in\s+([\d\.]+)\s*s', msg_erro, re.IGNORECASE)
                if not match_tempo:
                    match_tempo = re.search(r'retryDelay[\'\":\s]+(\d+)', msg_erro)
                tempo_espera = int(float(match_tempo.group(1))) + 3 if match_tempo else 20 * tentativa
                time.sleep(tempo_espera)
            else:
                raise err

    if not resposta:
        raise ultimo_erro

    return resposta.text, modelo_eleito

# ========================================================
# CRUZAMENTO INTELIGENTE E BIDIRECIONAL COM O PORTFÓLIO
# ========================================================
def enriquecer_com_portfolio(df_extraido, df_port, labs_escolhidos):
    if df_port is None:
        df_extraido["Laboratório Sugerido"] = "Sem base carregada"
        df_extraido["Produto / Marca Ref."] = "-"
        return df_extraido

    base = df_port.copy()
    base["SUBSTANCIA_NORM"] = base["SUBSTÂNCIA"].apply(normalizar_texto)
    base["LABORATORIO_NORM"] = base["LABORATÓRIO"].apply(normalizar_texto)
    base["PRODUTO_NORM"] = base["PRODUTO"].apply(normalizar_texto)

    termos_busca_labs = []
    for lab in labs_escolhidos:
        termos_busca_labs.extend(MAPEAMENTO_LABS.get(lab, [normalizar_texto(lab)]))

    mascara_apenas_medley = (base["LABORATORIO_NORM"] == "MEDLEY") & (~base["LABORATORIO_NORM"].str.contains("SANOFI"))
    base_valida = base[~mascara_apenas_medley]

    labs_atribuidos = []
    produtos_atribuidos = []

    for _, row in df_extraido.iterrows():
        substancia_edital = normalizar_texto(row["Princípio Ativo"])
        desc_completa = normalizar_texto(row["Descrição Completa Edital"])
        
        palavras_edital = [p for p in substancia_edital.split() if len(p) > 3]

        matches = base_valida[
            base_valida["SUBSTANCIA_NORM"].apply(lambda s: s != "" and (s in substancia_edital or substancia_edital in s or s in desc_completa)) |
            base_valida["PRODUTO_NORM"].apply(lambda p: p != "" and len(p) > 3 and (p in desc_completa or p in substancia_edital)) |
            base_valida["SUBSTANCIA_NORM"].apply(lambda s: any(p in s for p in palavras_edital[:2]) if len(palavras_edital) >= 1 else False)
        ]

        matches = matches[matches["LABORATORIO_NORM"].apply(
            lambda lab_nome: any(termo in lab_nome for termo in termos_busca_labs)
        )]

        if not matches.empty:
            labs_encontrados = matches["LABORATORIO_NORM"].unique().tolist()
            
            tem_sanofi = any("SANOFI" in l for l in labs_encontrados)
            tem_blau = any("BLAU" in l for l in labs_encontrados)
            tem_euro = any("EUROFARMA" in l for l in labs_encontrados)

            if tem_sanofi:
                lab_final_filtro = [l for l in labs_encontrados if "SANOFI" in l][0]
            elif tem_blau:
                lab_final_filtro = [l for l in labs_encontrados if "BLAU" in l][0]
            elif tem_euro:
                lab_final_filtro = [l for l in labs_encontrados if "EUROFARMA" in l][0]
            else:
                lab_final_filtro = labs_encontrados[0]

            linha_escolhida = matches[matches["LABORATORIO_NORM"] == lab_final_filtro].iloc[0]
            labs_atribuidos.append(linha_escolhida["LABORATÓRIO"])
            produtos_atribuidos.append(linha_escolhida["PRODUTO"])
        else:
            labs_atribuidos.append("Não mapeado / Verificar")
            produtos_atribuidos.append("-")

    df_extraido["Laboratório Sugerido"] = labs_atribuidos
    df_extraido["Produto / Marca Ref."] = produtos_atribuidos
    return df_extraido

# Gerador Excel
def gerar_excel_estilizado(df_dados):
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        df_dados.to_excel(writer, index=False, sheet_name="Mapa_de_Precos")
        ws = writer.sheets["Mapa_de_Precos"]
        ws.views.sheetView[0].showGridLines = True

        fonte_cabecalho = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        fill_cabecalho = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
        
        fonte_corpo = Font(name="Calibri", size=10)
        fill_editavel = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        
        alinhamento_centro = Alignment(horizontal="center", vertical="center")
        alinhamento_esquerda = Alignment(horizontal="left", vertical="center")
        alinhamento_direita = Alignment(horizontal="right", vertical="center")
        
        borda_fina = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9")
        )

        num_linhas = len(df_dados)
        num_cols = len(df_dados.columns)
        colunas_nomes = list(df_dados.columns)

        for col_num in range(1, num_cols + 1):
            celula = ws.cell(row=1, column=col_num)
            celula.font = fonte_cabecalho
            celula.fill = fill_cabecalho
            celula.alignment = alinhamento_centro
            ws.row_dimensions[1].height = 28

        def get_col_letter(col_name):
            if col_name in colunas_nomes:
                return get_column_letter(colunas_nomes.index(col_name) + 1)
            return None

        col_qtd = get_col_letter("Qtd")
        col_vref = get_col_letter("Valor Ref. Unit. (R$)")
        col_custo = get_col_letter("Custo Aquisição (R$)")
        col_margem = get_col_letter("Margem Alvo (%)")

        for row_idx in range(2, num_linhas + 2):
            ws.row_dimensions[row_idx].height = 20
            for col_idx in range(1, num_cols + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                nome_col = colunas_nomes[col_idx - 1]
                cell.font = fonte_corpo
                cell.border = borda_fina

                if nome_col in ["Item", "Unidade"]:
                    cell.alignment = alinhamento_centro
                elif nome_col in ["Descrição Completa Edital", "Princípio Ativo", "Laboratório Sugerido", "Produto / Marca Ref."]:
                    cell.alignment = alinhamento_esquerda
                elif nome_col == "Qtd":
                    cell.alignment = alinhamento_direita
                    cell.number_format = "#,##0"
                elif nome_col in ["Valor Ref. Unit. (R$)", "Valor Total Estimado (R$)"]:
                    cell.alignment = alinhamento_direita
                    cell.number_format = "R$ #,##0.00"
                    if nome_col == "Valor Total Estimado (R$)" and col_qtd and col_vref:
                        cell.value = f"={col_qtd}{row_idx}*{col_vref}{row_idx}"
                elif nome_col == "Custo Aquisição (R$)":
                    cell.alignment = alinhamento_direita
                    cell.number_format = "R$ #,##0.00"
                    cell.fill = fill_editavel
                elif nome_col == "Margem Alvo (%)":
                    cell.alignment = alinhamento_direita
                    cell.number_format = "0.0%"
                    cell.value = 0.15
                elif nome_col == "Preço Proposta Unit. (R$)":
                    cell.alignment = alinhamento_direita
                    cell.number_format = "R$ #,##0.00"
                    if col_custo and col_margem:
                        cell.value = f"={col_custo}{row_idx}*(1+{col_margem}{row_idx})"

        for col in ws.columns:
            col_letter = get_column_letter(col[0].column)
            max_len = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 13)

        if "Descrição Completa Edital" in colunas_nomes:
            ws.column_dimensions[get_col_letter("Descrição Completa Edital")].width = 42
        if "Princípio Ativo" in colunas_nomes:
            ws.column_dimensions[get_col_letter("Princípio Ativo")].width = 25
        if "Laboratório Sugerido" in colunas_nomes:
            ws.column_dimensions[get_col_letter("Laboratório Sugerido")].width = 28
        if "Produto / Marca Ref." in colunas_nomes:
            ws.column_dimensions[get_col_letter("Produto / Marca Ref.")].width = 22

    return output.getvalue()

# 5. Processamento com Barra de Progresso e Cache Instantâneo
if arquivo_pdf and api_key:
    if st.button("🚀 Processar Edital", type="primary"):
        barra_progresso = st.progress(0)
        texto_status = st.empty()
        
        try:
            pdf_bytes = arquivo_pdf.getvalue()
            
            # ETAPA 1: Leitura e Pré-processamento com Cache
            barra_progresso.progress(20)
            texto_status.text("📖 Verificando memória e filtrando páginas relevantes (20%)...")
            
            texto_edital, paginas_filtradas, total_pags = pre_processar_pdf_inteligente(pdf_bytes)

            if not texto_edital.strip():
                barra_progresso.empty()
                texto_status.empty()
                st.error("O PDF parece ser uma imagem digitalizada sem camada de texto pesquisável.")
                st.stop()

            # ETAPA 2: Preparação dos Metadados e Hash
            barra_progresso.progress(40)
            texto_status.text(f"⚙️ {len(paginas_filtradas)} páginas preparadas. Verificando cache de extração (40%)...")
            
            pdf_hash = hashlib.sha256((pdf_bytes + texto_edital.encode('utf-8'))).hexdigest()

            guia_substancias = ""
            if segmento == "Medicamentos" and df_portfolio is not None:
                subs_unicas = df_portfolio["SUBSTÂNCIA"].dropna().unique()[:250].tolist()
                guia_substancias = f"Lista de referência de substâncias prioritárias dos laboratórios parceiros:\n[{', '.join(subs_unicas)}]"

            if segmento == "Medicamentos":
                prompt = f"""
                Você é um analista sênior de licitações farmacêuticas e compras hospitalares.
                Analise o texto do edital e extraia ABSOLUTAMENTE TODOS os itens de medicamentos licitados.
                NÃO DEIXE NENHUM ITEM DE FORA.

                Laboratórios parceiros: [{", ".join(labs_selecionados)}]
                {guia_substancias}

                DIRETRIZES MANDATÓRIAS:
                1. Identifique e extraia todos os itens com descrição, dosagem e forma farmacêutica.
                2. No campo 'principio_ativo_identificado', extraia a denominação genérica exata da substância ativa (ex: 'Enoxaparina Sódica', 'Insulina Glargina', 'Propofol', 'Dipirona').
                3. Converta quantidades e valores de referência para números decimais (float).
                4. Retorne no formato JSON estruturado.

                TEXTO DO EDITAL:
                \"\"\"
                {texto_edital}
                \"\"\"
                """
            else:
                prompt = f"""
                Você é um analista de licitações. Extraia os itens correspondentes a: "{portfolio_texto}".
                Estruture no JSON solicitado com quantidades e valores unitários numéricos float.

                TEXTO DO EDITAL:
                \"\"\"
                {texto_edital}
                \"\"\"
                """

            # ETAPA 3: Extração Inteligente (Gemini ou Retorno Instantâneo de Cache)
            barra_progresso.progress(65)
            texto_status.text("🤖 Consultando IA / Cache de alta velocidade (65%)...")

            json_resposta, modelo_usado = extrair_itens_com_gemini_cached(
                pdf_hash, texto_edital, prompt, api_key
            )

            barra_progresso.progress(80)
            texto_status.text("📦 Dados carregados com sucesso. Estruturando linhas (80%)...")
            dados = ListaItens.model_validate_json(json_resposta)

            if not dados.itens:
                barra_progresso.empty()
                texto_status.empty()
                st.warning("Nenhum item foi identificado no edital.")
            else:
                # ETAPA 4: Cruzamento e Regras Comerciais
                barra_progresso.progress(90)
                texto_status.text("🧠 Aplicando hierarquia comercial (Sanofi > Blau > Eurofarma) (90%)...")

                lista_dicts = [item.model_dump() for item in dados.itens]
                salvar_aprendizado(lista_dicts)

                df = pd.DataFrame(lista_dicts)
                df.columns = [
                    "Item", 
                    "Descrição Completa Edital", 
                    "Princípio Ativo", 
                    "Unidade", 
                    "Qtd", 
                    "Valor Ref. Unit. (R$)"
                ]

                if (df["Valor Ref. Unit. (R$)"] > 0).any():
                    df["Valor Total Estimado (R$)"] = df["Qtd"] * df["Valor Ref. Unit. (R$)"]

                if segmento == "Medicamentos":
                    df = enriquecer_com_portfolio(df, df_portfolio, labs_selecionados)
                    if filtro_exibicao == "Apenas Itens com Match nos Laboratórios":
                        df = df[df["Laboratório Sugerido"] != "Não mapeado / Verificar"]

                df["Custo Aquisição (R$)"] = 0.0
                df["Margem Alvo (%)"] = 0.15
                df["Preço Proposta Unit. (R$)"] = 0.0

                try:
                    df["_item_num"] = pd.to_numeric(df["Item"].str.extract(r'(\d+)')[0], errors="coerce")
                    df = df.sort_values(by="_item_num").drop(columns=["_item_num"])
                except Exception:
                    pass

                # ETAPA 5: Construção do Excel Executivo
                barra_progresso.progress(96)
                texto_status.text("📊 Formatando planilha executiva com fórmulas (.xlsx) (96%)...")
                excel_bytes = gerar_excel_estilizado(df)

                # Finalização
                barra_progresso.progress(100)
                texto_status.text("✅ Processamento concluído!")
                time.sleep(0.3)
                barra_progresso.empty()
                texto_status.empty()

                st.success(f"✅ Mapeados **{len(df)} itens** no edital via **{modelo_usado}** (Cache Ativo)!")
                st.dataframe(df, use_container_width=True)

                st.download_button(
                    label="📥 Baixar Mapa de Preços Profissional (.xlsx)",
                    data=excel_bytes,
                    file_name=f"Mapa_Precos_{arquivo_pdf.name.replace('.pdf', '')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )

        except Exception as e:
            barra_progresso.empty()
            texto_status.empty()
            st.error(f"Erro no processamento: {str(e)}")

elif not api_key:
    st.info("Insira a sua chave de API na barra lateral ou configure nos secrets para iniciar a análise.")